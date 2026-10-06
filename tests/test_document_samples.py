from pathlib import Path

from docx import Document

from question_bank.parser import parse_docx_questions
from question_bank.parser import parse_preview_questions


def _create_synthetic_exam(path: Path) -> Path:
    document = Document()
    document.add_paragraph("单选题")
    for number in range(1, 20):
        document.add_paragraph(f"{number}. 计算 {number}+1 的值。")

    document.add_paragraph("参考答案与解析")
    for number in range(1, 20):
        document.add_paragraph(f"{number}. {number + 1}")
        document.add_paragraph("解析：根据加法运算计算。")

    document.save(path)
    return path


def test_synthetic_word_exam_finds_complete_19_question_sequence(tmp_path):
    sample = _create_synthetic_exam(tmp_path / "synthetic-exam.docx")

    questions = parse_docx_questions(sample, sample.name)

    assert [question.number for question in questions] == [str(i) for i in range(1, 20)]
    assert all(question.stem.strip() for question in questions)


def test_synthetic_concentrated_answers_are_merged(tmp_path):
    sample = _create_synthetic_exam(tmp_path / "synthetic-exam.docx")

    questions = parse_docx_questions(sample, sample.name)

    assert all(question.answer for question in questions)
    assert all(question.analysis for question in questions)
    assert questions[0].answer == "2"
    assert "加法运算" in questions[0].analysis

def test_preview_parser_preserves_question_tables(tmp_path):
    preview = tmp_path / "preview.htm"
    preview.write_text(
        """
        <html><body>
          <p>1．某题原始分及转换分如下表：</p>
          <table><tr><td><p>原始分</p></td><td><p>91</p></td></tr><tr><td><p>转换分</p></td><td><p>100</p></td></tr></table>
          <p>现从这10名学生中随机抽取3人，求分布列。</p>
        </body></html>
        """,
        encoding="utf-8",
    )
    questions = parse_preview_questions(preview, source_name="preview.docx", data_url_prefix="/data/previews/sample")
    assert len(questions) == 1
    assert "<table>" in questions[0].stem_html
    assert "原始分" in questions[0].stem

def test_preview_parser_strips_split_span_question_number(tmp_path):
    preview = tmp_path / "preview.htm"
    preview.write_text(
        """
        <html><body>
          <p><span>19</span><span>．某省新高考方案。</span></p>
        </body></html>
        """,
        encoding="utf-8",
    )
    questions = parse_preview_questions(preview, source_name="preview.docx", data_url_prefix="/data/previews/sample")
    assert len(questions) == 1
    assert questions[0].stem.startswith("某省")
    assert ">19<" not in questions[0].stem_html
