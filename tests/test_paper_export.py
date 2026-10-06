import io,zipfile
from docx import Document
from question_bank.papers import build_export_model
from question_bank.exporter import export_paper_word_zip,export_paper_pdf

def _paper():
    return {"title":"期中测试","exam_time":"120 分钟","total_score":"150","school":"示例学校","grade":"高二","class_name":"1班","instructions":"请认真答题。","questions":[
        {"id":3,"position":1,"question_type":"解答题","stem":"解答题干","options":"","answer":"结论","analysis":"过程"},
        {"id":1,"position":2,"question_type":"单选题","stem":"选择题干","options":"A. 1\nB. 2","answer":"A","analysis":"解析"},
        {"id":2,"position":3,"question_type":"填空题","stem":"填空题干","options":"","answer":"3","analysis":"解析"}]}

def test_export_model_groups_and_renumbers():
    model=build_export_model(_paper())
    assert [q["question_type"] for q in model.questions]==["单选题","填空题","解答题"]
    assert [q["export_number"] for q in model.questions]==[1,2,3]

def test_word_zip_has_exactly_two_valid_documents(tmp_path):
    path=export_paper_word_zip(build_export_model(_paper()),tmp_path/"paper.zip")
    with zipfile.ZipFile(path) as z:
        assert sorted(z.namelist())==["试卷-答案与解析.docx","试卷-题目.docx"]
        question=Document(io.BytesIO(z.read("试卷-题目.docx"))); answer=Document(io.BytesIO(z.read("试卷-答案与解析.docx")))
    qt="\n".join(p.text for p in question.paragraphs); at="\n".join(p.text for p in answer.paragraphs)
    assert "考试时间：120 分钟" in qt and "答案：" not in qt
    assert "选择题干" in at and "答案：" in at and "考试时间" not in at

def test_real_pdf_export(tmp_path):
    path=export_paper_pdf(build_export_model(_paper()),tmp_path/"paper.pdf")
    assert path.read_bytes().startswith(b"%PDF-") and path.stat().st_size>1000


def test_pdf_export_falls_back_when_office_conversion_fails(tmp_path, monkeypatch):
    def fail_conversion(_docx_path, _output_path):
        raise RuntimeError("office conversion failed")

    monkeypatch.setattr("question_bank.exporter._convert_docx_to_pdf_with_word", fail_conversion)
    path=export_paper_pdf(build_export_model(_paper()),tmp_path/"fallback.pdf")
    assert path.read_bytes().startswith(b"%PDF") and path.stat().st_size>1000
