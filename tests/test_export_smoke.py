from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook
from PIL import Image


def test_xlsx_export_route(client, app, seeded_question):
    response = client.get("/export/xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.active.title == "题库"
    assert workbook.active.max_row == 2


def test_docx_export_route(client, app, seeded_question):
    response = client.get("/export/docx")
    assert response.status_code == 200
    document = Document(io.BytesIO(response.data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "高中数学题库导出" in text
    assert "已知集合 A" in text
    assert "答案" not in text
    assert "一、单选题" in text


def test_grouped_docx_export_omits_metadata_labels(tmp_path):
    from question_bank.exporter import export_questions_to_docx

    path = export_questions_to_docx(
        [
            {"number": "9", "question_type": "多选题", "category_major": "概率统计", "category_minor": "概率", "difficulty": "困难", "source_name": "卷A", "stem": "多选题内容", "options": "A. 1\nB. 2", "answer": "AB", "analysis": "解析"},
            {"number": "1", "question_type": "单选题", "category_major": "集合", "category_minor": "集合运算", "difficulty": "简单", "source_name": "卷B", "stem": "单选题内容", "options": "A. 1\nB. 2", "answer": "A", "analysis": "解析"},
        ],
        tmp_path / "grouped.docx",
    )
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "一、单选题" in text
    assert "二、多选题" in text
    assert "[单选题" not in text
    assert "概率统计" not in text
    assert "卷A" not in text


def test_stems_only_export_has_no_options_or_answers(tmp_path):
    from question_bank.exporter import export_questions_to_docx

    path = export_questions_to_docx(
        [{"number": "1", "question_type": "单选题", "stem": "题干内容", "options": "A. 选项", "answer": "A", "analysis": "解析"}],
        tmp_path / "stems.docx",
        stems_only=True,
    )
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "题干内容" in text
    assert "选项" not in text
    assert "答案" not in text
    assert "解析" not in text


def test_docx_export_preserves_html_tables(tmp_path):
    from question_bank.exporter import export_questions_to_docx

    path = export_questions_to_docx(
        [
            {
                "number": "1",
                "question_type": "解答题",
                "stem_html": "<p>原始分及转换分如下表：</p><table><tr><td>原始分</td><td>91</td></tr><tr><td>转换分</td><td>100</td></tr></table>",
                "stem": "原始分及转换分如下表：",
                "options": "",
            }
        ],
        tmp_path / "table.docx",
    )
    document = Document(path)
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "原始分"
    assert document.tables[0].cell(1, 1).text == "100"


def test_docx_export_restores_flattened_score_tables(tmp_path):
    from question_bank.exporter import export_questions_to_docx

    path = export_questions_to_docx(
        [
            {
                "number": "19",
                "question_type": "解答题",
                "stem_html": (
                    "<p>某校学生成绩如下表：</p>"
                    "<p>原始分</p><p>91</p><p>90</p>"
                    "<p>转换分</p><p>100</p><p>99</p>"
                    "<p>人数</p><p>1</p><p>2</p>"
                    "<p>现从这10名学生中随机抽取3人。</p>"
                ),
                "stem": "某校学生成绩如下表：",
                "options": "",
            }
        ],
        tmp_path / "restored-table.docx",
    )
    document = Document(path)
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "原始分"
    assert document.tables[0].cell(0, 2).text == "90"
    assert document.tables[0].cell(2, 2).text == "2"


def test_docx_export_enlarges_block_figures(tmp_path):
    from question_bank.exporter import export_questions_to_docx

    image_path = tmp_path / "figure.png"
    Image.new("RGB", (223, 130), "white").save(image_path)
    path = export_questions_to_docx(
        [
            {
                "number": "1",
                "question_type": "单选题",
                "stem_html": f'<p>函数图像如下：</p><p><img class="inline-math block-figure" src="{image_path}"></p>',
                "stem": "函数图像如下：",
                "options": "",
            }
        ],
        tmp_path / "figure.docx",
    )
    document = Document(path)
    assert document.inline_shapes
    assert document.inline_shapes[0].width >= Inches(2.3)


def test_image_zip_export_route(client, app, seeded_question):
    response = client.get("/export/images")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        names = archive.namelist()
        assert len(names) == 1
        assert names[0].endswith(".png")


def test_pdf_export_route_can_dispatch(client, app, seeded_question, monkeypatch):
    def fake_pdf_export(_questions, output_path):
        output_path.write_bytes(b"%PDF-1.4\n% stage-zero smoke\n")
        return output_path

    monkeypatch.setattr("question_bank.app.export_questions_to_pdf", fake_pdf_export)
    response = client.get("/export/pdf")
    assert response.status_code == 200
    assert response.data.startswith(b"%PDF-1.4")
