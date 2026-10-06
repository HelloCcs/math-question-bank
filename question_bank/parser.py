from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass, replace
from html import escape
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urlparse

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from PIL import Image

QUESTION_RE = re.compile(
    r"^\s*(?:[\(\uFF08\u3008\u3010]\s*[^\(\)\uFF08\uFF09\u3008\u3009\u3010\u30110-9]{1,6}\s*[\)\uFF09\u3009\u3011]\s*)*"
    r"(?:[\(\uFF08\u3008\u3010]\s*(?P<bracketed>[0-9\uFF10-\uFF19]{1,3})\s*[\)\uFF09\u3009\u3011]|"
    r"(?P<plain>[0-9\uFF10-\uFF19]{1,3})\s*(?:[\.\uFF0E\u3001])\s*(?P<body>.*)$)"
)
OPTION_RE = re.compile(r"^\s*([A-HＡ-Ｈ])\s*[\.．、]\s*(.+)$", re.IGNORECASE)
ANSWER_RE = re.compile(r"^\s*[【\[]?\s*(?:答案|参考答案|正确答案)\s*[】\]]?\s*[:：]?\s*(.*)$")
ANALYSIS_RE = re.compile(r"^\s*[【\[]?\s*(?:解析|解答|详解|分析)\s*[】\]]?\s*[:：]?\s*(.*)$")
SECTION_KEYWORDS = (
    "单选题", "选择题", "多选题", "填空题", "解答题", "计算题", "证明题", 
    "问答题", "判断题", "判断题", "计算题", "应用题", "探究题", "开放题"
)

SOLUTION_KEYWORDS = ("证明", "求证", "求解", "计算", "化简", "求值", "求导", 
                     "讨论", "分析", "说明", "推导", "证明如下", "解答", "解题")
SECTION_PREFIX_RE = re.compile(
    r"^(?:[\u4e00-\u4e5d\u5341\u767e]+[\u3001\uff1a.:]|[0-9]+[\u3001\uff1a.:]|\u7b2c[\u4e00-\u4e5d\u53410-9]+[\u5927\u90e8\u5377\u7ae0\u8282])"
)


@dataclass(frozen=True)
class QuestionDraft:
    number: str
    stem: str
    options: list[str]
    answer: str
    analysis: str
    question_type: str
    source_name: str
    raw_text: str
    stem_html: str = ""
    options_html: str = ""
    answer_html: str = ""
    analysis_html: str = ""
    raw_html: str = ""
    category_major: str = ""
    category_minor: str = ""
    difficulty: str = ""
    curriculum_version: str = "人教B版"
    curriculum_chapter: str = ""
    curriculum_section: str = ""
    classification_source: str = "rule"
    classification_confidence: float | None = None
    review_reason: str = ""
    rule_category_major: str = ""
    rule_category_minor: str = ""
    ai_category_major: str = ""
    ai_category_minor: str = ""
    ai_review_note: str = ""
    asset_manifest: list[dict] | None = None

    def __post_init__(self) -> None:
        if self.asset_manifest is None:
            object.__setattr__(self, "asset_manifest", [])


@dataclass(frozen=True)
class AssetSummary:
    filename: str
    media_counts: dict[str, int]
    formula_count: int
    page_count: int | None


def analyze_docx_assets(path: str | Path) -> AssetSummary:
    docx_path = Path(path)
    media_counts: dict[str, int] = {}
    formula_count = 0
    page_count: int | None = None

    with zipfile.ZipFile(docx_path) as archive:
        for entry in archive.infolist():
            if not entry.filename.startswith("word/media/") or entry.is_dir():
                continue
            ext = Path(entry.filename).suffix.lower().lstrip(".") or "no_extension"
            media_counts[ext] = media_counts.get(ext, 0) + 1

        try:
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
            formula_count = len(re.findall(r"<m:oMath(?:\s|>)", document_xml))
        except KeyError:
            formula_count = 0

        try:
            app_xml = archive.read("docProps/app.xml").decode("utf-8", errors="ignore")
            match = re.search(r"<Pages>(\d+)</Pages>", app_xml)
            if match:
                page_count = int(match.group(1))
        except KeyError:
            page_count = None

    return AssetSummary(
        filename=docx_path.name,
        media_counts=dict(sorted(media_counts.items())),
        formula_count=formula_count,
        page_count=page_count,
    )


def parse_docx_questions(path: str | Path, source_name: str | None = None) -> list[QuestionDraft]:
    docx_path = Path(path)
    display_name = source_name or docx_path.name
    document = Document(docx_path)
    questions: list[QuestionDraft] = []
    current_type = ""
    current_question_type = ""
    current_number: str | None = None
    current_lines: list[str] = []
    answer_section_lines: list[str] = []
    in_answer_section = False

    for raw_line in _iter_paragraph_text(document):
        if _is_answer_section_heading(raw_line):
            if current_number is not None and current_lines:
                questions.append(_build_question(current_number, current_lines, current_question_type, display_name))
                current_number = None
                current_lines = []
            in_answer_section = True
            continue
        if in_answer_section:
            cleaned = _clean_line(raw_line)
            if cleaned:
                answer_section_lines.append(cleaned)
            continue
        for line in _split_possible_question_lines(_clean_line(raw_line)):
            if not line:
                continue

            section_type = _detect_section_type(line)
            if section_type and not QUESTION_RE.match(line):
                if current_number is not None and current_lines:
                    questions.append(
                        _build_question(current_number, current_lines, current_question_type, display_name)
                    )
                    current_number = None
                    current_lines = []
                current_type = section_type
                continue

            question_match = QUESTION_RE.match(line)
            if question_match:
                candidate_number = _normalize_digits(_question_number(question_match))
                if not _should_start_new_question(current_number, candidate_number):
                    if current_number is not None:
                        current_lines.append(line)
                    continue
                if current_number is not None and current_lines:
                    questions.append(
                        _build_question(current_number, current_lines, current_question_type, display_name)
                    )
                current_number = candidate_number
                current_question_type = current_type
                first_line = _question_body(question_match).strip()
                current_lines = [first_line] if first_line else []
                continue

            if current_number is not None:
                current_lines.append(line)

    if current_number is not None and current_lines:
        questions.append(
            _build_question(current_number, current_lines, current_question_type, display_name)
        )

    questions = _drop_preface_questions(questions)
    return _merge_answer_section(questions, answer_section_lines)


def parse_docx_single_question(path: str | Path, source_name: str | None = None) -> list[QuestionDraft]:
    docx_path = Path(path)
    display_name = source_name or docx_path.name
    document = Document(docx_path)
    current_type = ""
    lines: list[str] = []
    number = "1"
    first_content_seen = False

    for raw_line in _iter_paragraph_text(document):
        cleaned = _clean_line(raw_line)
        if not cleaned:
            continue
        if _is_answer_section_heading(cleaned):
            break
        section_type = _detect_section_type(cleaned)
        if section_type and not QUESTION_RE.match(cleaned):
            current_type = section_type
            continue

        if not first_content_seen:
            question_match = QUESTION_RE.match(cleaned)
            if question_match:
                candidate_number = _normalize_digits(_question_number(question_match))
                if candidate_number:
                    number = candidate_number
                cleaned = _question_body(question_match).strip()
            first_content_seen = True

        if cleaned:
            lines.append(cleaned)

    if not lines:
        return []
    return [_build_question(number, lines, current_type, display_name)]


def parse_plain_text_questions(text: str, source_name: str) -> list[QuestionDraft]:
    questions: list[QuestionDraft] = []
    current_type = ""
    current_question_type = ""
    current_number: str | None = None
    current_lines: list[str] = []

    for raw_line in text.splitlines():
        if _is_answer_section_heading(raw_line):
            break
        for line in _split_possible_question_lines(_clean_line(raw_line)):
            if not line:
                continue

            section_type = _detect_section_type(line)
            if section_type and not QUESTION_RE.match(line):
                if current_number is not None and current_lines:
                    questions.append(_build_question(current_number, current_lines, current_question_type, source_name))
                    current_number = None
                    current_lines = []
                current_type = section_type
                continue

            question_match = QUESTION_RE.match(line)
            if question_match:
                candidate_number = _normalize_digits(_question_number(question_match))
                if not _should_start_new_question(current_number, candidate_number):
                    if current_number is not None:
                        current_lines.append(line)
                    continue
                if current_number is not None and current_lines:
                    questions.append(_build_question(current_number, current_lines, current_question_type, source_name))
                current_number = candidate_number
                current_question_type = current_type
                first_line = _question_body(question_match).strip()
                current_lines = [first_line] if first_line else []
                continue

            if current_number is not None:
                current_lines.append(line)

    if current_number is not None and current_lines:
        questions.append(_build_question(current_number, current_lines, current_question_type, source_name))

    questions = _drop_preface_questions(questions)
    if questions:
        return questions
    cleaned = "\n".join(_clean_line(line) for line in text.splitlines() if _clean_line(line))
    if not cleaned:
        return []
    return [
        QuestionDraft(
            number="1",
            stem=cleaned,
            options=[],
            answer="",
            analysis="",
            question_type="OCR识别",
            source_name=source_name,
            raw_text=cleaned,
            stem_html=f"<p>{escape(cleaned)}</p>",
            raw_html=f"<p>{escape(cleaned)}</p>",
        )
    ]


def parse_preview_questions(
    preview_path: str | Path,
    *,
    source_name: str,
    data_url_prefix: str,
) -> list[QuestionDraft]:
    html_path = Path(preview_path)
    soup = BeautifulSoup(_read_html(html_path), "html.parser")
    questions: list[QuestionDraft] = []
    current_type = ""
    current_question_type = ""
    current_number: str | None = None
    current_lines: list[tuple[str, str]] = []
    answer_section_lines: list[tuple[str, str]] = []
    in_answer_section = False

    for block in _iter_preview_blocks(soup):
        line_text = _clean_line(block.get_text(" ", strip=True))
        line_html = _block_to_html(block, data_url_prefix, html_path.parent)
        if not line_text and "<img" not in line_html:
            continue

        if block.name == "table":
            if in_answer_section:
                answer_section_lines.append((line_text, line_html))
            elif current_number is not None:
                current_lines.append((line_text, line_html))
            continue

        if _is_answer_section_heading(line_text):
            if current_number is not None and current_lines:
                questions.append(
                    _build_question_from_rich_lines(
                        current_number, current_lines, current_question_type, source_name
                    )
                )
                current_number = None
                current_lines = []
            in_answer_section = True
            continue

        if in_answer_section:
            answer_section_lines.append((line_text, line_html))
            continue

        for segment_index, segment_text in enumerate(_split_possible_question_lines(line_text)):
            segment_html = line_html if segment_index == 0 else escape(segment_text)
            section_type = _detect_section_type(segment_text)
            if section_type and not QUESTION_RE.match(segment_text):
                if current_number is not None and current_lines:
                    questions.append(
                        _build_question_from_rich_lines(
                            current_number,
                            current_lines,
                            current_question_type,
                            source_name,
                        )
                    )
                    current_number = None
                    current_lines = []
                current_type = section_type
                continue

            question_match = QUESTION_RE.match(segment_text)
            if question_match:
                candidate_number = _normalize_digits(_question_number(question_match))
                if not _should_start_new_question(current_number, candidate_number):
                    if current_number is not None:
                        current_lines.append((segment_text, segment_html))
                    continue
                if current_number is not None and current_lines:
                    questions.append(
                        _build_question_from_rich_lines(
                            current_number,
                            current_lines,
                            current_question_type,
                            source_name,
                        )
                    )
                current_number = candidate_number
                current_question_type = current_type
                segment_text = _question_body(question_match).strip()
                segment_html = _strip_question_number_from_html(segment_html, _question_number(question_match))
                current_lines = [(segment_text, segment_html)] if segment_text or segment_html else []
                continue

            if current_number is not None:
                current_lines.append((segment_text, segment_html))

    if current_number is not None and current_lines:
        questions.append(
            _build_question_from_rich_lines(
                current_number,
                current_lines,
                current_question_type,
                source_name,
            )
        )

    questions = _drop_preface_questions(questions)
    return _merge_rich_answer_section(questions, answer_section_lines)


def parse_preview_single_question(
    preview_path: str | Path,
    *,
    source_name: str,
    data_url_prefix: str,
) -> list[QuestionDraft]:
    html_path = Path(preview_path)
    soup = BeautifulSoup(_read_html(html_path), "html.parser")
    current_type = ""
    lines: list[tuple[str, str]] = []
    number = "1"
    first_content_seen = False

    for block in _iter_preview_blocks(soup):
        line_text = _clean_line(block.get_text(" ", strip=True))
        line_html = _block_to_html(block, data_url_prefix, html_path.parent)
        if not line_text and "<img" not in line_html:
            continue

        if block.name == "table":
            if first_content_seen:
                lines.append((line_text, line_html))
            continue

        section_type = _detect_section_type(line_text)
        if section_type and not QUESTION_RE.match(line_text):
            current_type = section_type
            continue

        if not first_content_seen:
            question_match = QUESTION_RE.match(line_text)
            if question_match:
                candidate_number = _normalize_digits(_question_number(question_match))
                if candidate_number:
                    number = candidate_number
                line_text = _question_body(question_match).strip()
                line_html = _strip_question_number_from_html(line_html, _question_number(question_match))
            first_content_seen = True

        if line_text or line_html:
            lines.append((line_text, line_html))

    if not lines:
        return []
    return [_build_question_from_rich_lines(number, lines, current_type, source_name)]


def _iter_paragraph_text(document: Document) -> Iterable[str]:
    for paragraph in document.paragraphs:
        parts: list[str] = []
        for node in paragraph._p.iter():
            tag = node.tag.rsplit("}", 1)[-1]
            if tag == "t" and node.text:
                parts.append(node.text)
            elif tag == "tab":
                parts.append("\t")
            elif tag == "br":
                parts.append("\n")
        text = "".join(parts).strip()
        if text:
            yield text


def _build_question(
    number: str,
    lines: list[str],
    current_type: str,
    source_name: str,
) -> QuestionDraft:
    stem_lines: list[str] = []
    options: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []
    mode = "stem"

    for source_line in lines:
        for line in _split_inline_answer_analysis(source_line):
            answer_match = ANSWER_RE.match(line)
            if answer_match:
                mode = "answer"
                if answer_match.group(1).strip():
                    answer_lines.append(answer_match.group(1).strip())
                continue

            analysis_match = ANALYSIS_RE.match(line)
            if analysis_match:
                mode = "analysis"
                if analysis_match.group(1).strip():
                    analysis_lines.append(analysis_match.group(1).strip())
                continue

            if mode == "answer":
                answer_lines.append(line)
                continue
            if mode == "analysis":
                analysis_lines.append(line)
                continue

            option_match = OPTION_RE.match(line)
            if option_match:
                option_label = option_match.group(1).upper()
                options.append(f"{option_label}. {option_match.group(2).strip()}")
            else:
                stem_lines.append(line)

    answer_text = "\n".join(answer_lines).strip()
    stem_text = "\n".join(stem_lines).strip()
    
    if current_type:
        question_type = current_type
    elif _looks_like_solution(stem_text):
        question_type = "解答题"
    elif _looks_like_solution(answer_text):
        question_type = "解答题"
    elif options and len(options) <= 6:
        question_type = "单选题"
    elif options and len(options) > 6:
        question_type = "解答题"
    elif re.fullmatch(r"[A-HＡ-Ｈ](?:\s*[,，、]\s*[A-HＡ-Ｈ])+", answer_text, re.IGNORECASE):
        if "," in answer_text or "、" in answer_text:
            question_type = "多选题"
        else:
            question_type = "单选题"
    elif re.fullmatch(r"[A-HＡ-Ｈ]", answer_text, re.IGNORECASE):
        question_type = "单选题"
    elif answer_text and len(answer_text) <= 15 and not _looks_like_sentence(answer_text):
        question_type = "填空题"
    else:
        question_type = "解答题"
    raw_text = "\n".join(lines).strip()
    return QuestionDraft(
        number=number,
        stem=stem_text or raw_text,
        options=options,
        answer=answer_text,
        analysis="\n".join(analysis_lines).strip(),
        question_type=question_type,
        source_name=source_name,
        raw_text=raw_text,
    )


def _build_question_from_rich_lines(
    number: str,
    lines: list[tuple[str, str]],
    current_type: str,
    source_name: str,
) -> QuestionDraft:
    stem_lines: list[str] = []
    options: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []
    stem_html: list[str] = []
    options_html: list[str] = []
    answer_html: list[str] = []
    analysis_html: list[str] = []
    mode = "stem"

    for source_line, source_html in lines:
        rich_segments = _split_inline_rich_segments(source_line, source_html)
        for line, line_html in rich_segments:
            answer_match = ANSWER_RE.match(line)
            if answer_match:
                mode = "answer"
                answer_text = answer_match.group(1).strip()
                answer_line_html = _strip_label_from_html(
                    line_html, _label_from_match(answer_match)
                )
                if answer_text:
                    answer_lines.append(answer_text)
                if answer_line_html:
                    answer_html.append(answer_line_html)
                continue

            analysis_match = ANALYSIS_RE.match(line)
            if analysis_match:
                mode = "analysis"
                analysis_text = analysis_match.group(1).strip()
                analysis_line_html = _strip_label_from_html(
                    line_html, _label_from_match(analysis_match)
                )
                if analysis_text:
                    analysis_lines.append(analysis_text)
                if analysis_line_html:
                    analysis_html.append(analysis_line_html)
                continue

            if mode == "answer":
                answer_lines.append(line)
                answer_html.append(line_html)
                continue
            if mode == "analysis":
                analysis_lines.append(line)
                analysis_html.append(line_html)
                continue

            option_match = OPTION_RE.match(line)
            if option_match:
                option_label = option_match.group(1).upper()
                options.append(f"{option_label}. {option_match.group(2).strip()}")
                options_html.append(line_html)
            elif _looks_like_option_group(line):
                options.append(line)
                options_html.append(line_html)
            else:
                stem_lines.append(line)
                stem_html.append(line_html)

    answer_text = "\n".join(answer_lines).strip()
    stem_text = "\n".join(stem_lines).strip()
    
    if current_type:
        question_type = current_type
    elif _looks_like_solution(stem_text):
        question_type = "解答题"
    elif _looks_like_solution(answer_text):
        question_type = "解答题"
    elif options and len(options) <= 6:
        question_type = "单选题"
    elif options and len(options) > 6:
        question_type = "解答题"
    elif re.fullmatch(r"[A-HＡ-Ｈ](?:\s*[,，、]\s*[A-HＡ-Ｈ])+", answer_text, re.IGNORECASE):
        if "," in answer_text or "、" in answer_text:
            question_type = "多选题"
        else:
            question_type = "单选题"
    elif re.fullmatch(r"[A-HＡ-Ｈ]", answer_text, re.IGNORECASE):
        question_type = "单选题"
    elif answer_text and len(answer_text) <= 15 and not _looks_like_sentence(answer_text):
        question_type = "填空题"
    else:
        question_type = "解答题"
    raw_text = "\n".join(line for line, _ in lines).strip()
    raw_html = "".join(_wrap_html_line(line_html) for _, line_html in lines if line_html)
    return QuestionDraft(
        number=number,
        stem="\n".join(stem_lines).strip() or raw_text,
        options=options,
        answer=answer_text,
        analysis="\n".join(analysis_lines).strip(),
        question_type=question_type,
        source_name=source_name,
        raw_text=raw_text,
        stem_html=_join_html_lines(stem_html),
        options_html=_join_html_lines(options_html),
        answer_html=_join_html_lines(answer_html),
        analysis_html=_join_html_lines(analysis_html),
        raw_html=raw_html,
    )


def _clean_line(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u3000", " ")).strip()


def _split_possible_question_lines(line: str) -> list[str]:
    if not line:
        return [line]
    return _split_inline_questions_safely(line)


INLINE_QUESTION_RE = re.compile(
    r"\s+(?=(?:[\(\uFF08\u3008\u3010]\s*[0-9\uFF10-\uFF19]{1,3}\s*[\)\uFF09\u3009\u3011]|"
    r"[0-9\uFF10-\uFF19]{1,3}\s*(?:[\.\uFF0E\u3001](?![0-9\uFF10-\uFF19])))\s*\S)"
)


def _split_inline_questions_safely(line: str) -> list[str]:
    segments: list[str] = []
    remaining = line.strip()
    while remaining:
        boundary = _find_inline_question_boundary(remaining)
        if boundary is None:
            segments.append(remaining.strip())
            break
        head = remaining[:boundary].strip()
        if head:
            segments.append(head)
        remaining = remaining[boundary:].strip()
    return segments or [line]


def _find_inline_question_boundary(line: str) -> int | None:
    for match in INLINE_QUESTION_RE.finditer(line):
        candidate = line[match.end():].lstrip()
        if not QUESTION_RE.match(candidate):
            continue
        prefix = line[:match.start()]
        if _looks_like_question_boundary(prefix):
            return match.end()
    return None


def _looks_like_question_boundary(prefix: str) -> bool:
    compact = prefix.strip()
    if not compact:
        return False
    tail = compact[-80:]
    labels = (
        "答案",
        "解析",
        "详解",
        "分析",
        "参考答案",
        "故选",
        "故答案为",
    )
    if any(label in tail for label in labels):
        return True
    if len(re.findall(r"(?:^|\s)[A-H]\s*[\.\uFF0E\u3001]", tail, re.IGNORECASE)) >= 2:
        return True
    if re.search(r"[\u3002\uFF01\uFF1F\uFF1B;:?]\s*$", compact):
        return True
    return False


def _question_number(match: re.Match[str]) -> str:
    return match.group("bracketed") or match.group("plain") or ""


def _question_body(match: re.Match[str]) -> str:
    return match.group("body") or ""


def _is_answer_section_heading(line: str) -> bool:
    cleaned = _clean_line(line)
    return "参考答案" in cleaned or "评分标准" in cleaned or bool(re.search(r"数学答案\s*$", cleaned))


ANSWER_ENTRY_RE = re.compile(r"^\s*([0-9\uFF10-\uFF19]{1,3})\s*[\.\uFF0E\u3001]\s*(.*)$")


def _merge_answer_section(questions: list[QuestionDraft], lines: list[str]) -> list[QuestionDraft]:
    if not questions or not lines:
        return questions
    entries: dict[str, list[str]] = {}
    current = ""
    for line in lines:
        match = ANSWER_ENTRY_RE.match(line)
        if match:
            number = _normalize_digits(match.group(1))
            if number in {question.number for question in questions}:
                current = number
                entries.setdefault(current, [])
                body = match.group(2).strip()
                if body:
                    entries[current].append(body)
                continue
        if current:
            entries[current].append(line)

    merged: list[QuestionDraft] = []
    for question in questions:
        entry = entries.get(question.number)
        if not entry:
            merged.append(question)
            continue
        answer, analysis = _answer_entry_content(entry)
        merged.append(
            replace(
                question,
                answer=question.answer or answer,
                analysis=question.analysis or analysis,
                answer_html=question.answer_html or (f"<p>{escape(answer)}</p>" if answer else ""),
                analysis_html=question.analysis_html or (f"<p>{escape(analysis)}</p>" if analysis else ""),
            )
        )
    return merged


def _merge_rich_answer_section(
    questions: list[QuestionDraft], lines: list[tuple[str, str]]
) -> list[QuestionDraft]:
    if not questions or not lines:
        return questions
    valid_numbers = {question.number for question in questions}
    entries: dict[str, list[tuple[str, str]]] = {}
    current = ""
    for line_text, line_html in lines:
        match = ANSWER_ENTRY_RE.match(line_text)
        if match:
            number = _normalize_digits(match.group(1))
            if number in valid_numbers:
                current = number
                entries.setdefault(current, [])
                body_text = match.group(2).strip()
                body_html = _strip_question_number_from_html(
                    line_html, match.group(1)
                )
                if body_text or body_html:
                    entries[current].append((body_text, body_html))
                continue
        if current:
            entries[current].append((line_text, line_html))

    merged: list[QuestionDraft] = []
    for question in questions:
        entry = entries.get(question.number)
        if not entry:
            merged.append(question)
            continue
        answer, analysis, answer_html, analysis_html = _rich_answer_entry_content(entry)
        merged.append(
            replace(
                question,
                answer=question.answer or answer,
                analysis=question.analysis or analysis,
                answer_html=question.answer_html or answer_html,
                analysis_html=question.analysis_html or analysis_html,
            )
        )
    return merged


def _rich_answer_entry_content(
    lines: list[tuple[str, str]]
) -> tuple[str, str, str, str]:
    answer_text: list[str] = []
    analysis_text: list[str] = []
    answer_html: list[str] = []
    analysis_html: list[str] = []
    mode = "answer"
    for line, line_html in lines:
        answer_match = ANSWER_RE.match(line)
        analysis_match = ANALYSIS_RE.match(line)
        if answer_match:
            mode = "answer"
            value = answer_match.group(1).strip()
            if value:
                answer_text.append(value)
            stripped = _strip_label_from_html(line_html, _label_from_match(answer_match))
            if stripped:
                answer_html.append(stripped)
            continue
        if analysis_match:
            mode = "analysis"
            value = analysis_match.group(1).strip()
            if value:
                analysis_text.append(value)
            stripped = _strip_label_from_html(line_html, _label_from_match(analysis_match))
            if stripped:
                analysis_html.append(stripped)
            continue
        if mode == "analysis":
            analysis_text.append(line)
            analysis_html.append(line_html)
        else:
            answer_text.append(line)
            answer_html.append(line_html)
    return (
        "\n".join(answer_text).strip(),
        "\n".join(analysis_text).strip(),
        _join_html_lines(answer_html),
        _join_html_lines(analysis_html),
    )


def _label_from_match(match: re.Match[str]) -> str:
    value = match.group(0)
    for label in ("参考答案", "正确答案", "答案", "解析", "解答", "详解", "分析"):
        if label in value:
            return label
    return ""


def _answer_entry_content(lines: list[str]) -> tuple[str, str]:
    answers: list[str] = []
    analyses: list[str] = []
    mode = "answer"
    for source in lines:
        for line in _split_inline_answer_analysis(source):
            answer_match = ANSWER_RE.match(line)
            if answer_match:
                mode = "answer"
                if answer_match.group(1).strip():
                    answers.append(answer_match.group(1).strip())
                continue
            analysis_match = ANALYSIS_RE.match(line)
            if analysis_match:
                mode = "analysis"
                if analysis_match.group(1).strip():
                    analyses.append(analysis_match.group(1).strip())
                continue
            (analyses if mode == "analysis" else answers).append(line)
    return "\n".join(answers).strip(), "\n".join(analyses).strip()


INLINE_LABEL_RE = re.compile(
    r"(\s*(?:[\u3010\[]\s*(?:答案|参考答案|正确答案|解析|解答|详解|分析)\s*[\u3011\]]|(?:答案|参考答案|正确答案|解析|解答|详解|分析)\s*[:：]))"
)


def _split_inline_answer_analysis(line: str) -> list[str]:
    if not line or ANSWER_RE.match(line) or ANALYSIS_RE.match(line):
        return [line]
    matches = list(INLINE_LABEL_RE.finditer(line))
    if not matches:
        return [line]

    segments: list[str] = []
    if matches[0].start() > 0:
        segments.append(line[: matches[0].start()].strip())
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
        segments.append(line[match.start() : end].strip())
    return [segment for segment in segments if segment]


def _split_inline_rich_segments(line: str, line_html: str) -> list[tuple[str, str]]:
    segments = _split_inline_answer_analysis(line)
    if len(segments) == 1:
        return [(line, line_html)]
    return [(segment, escape(segment)) for segment in segments]


def _read_html(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _iter_preview_blocks(soup: BeautifulSoup):
    tags = ["p", "h1", "h2", "h3", "h4", "h5", "h6", "table"]
    root = soup.body or soup
    for block in root.find_all(tags):
        if block.find_parent("table"):
            continue
        yield block


def _block_to_html(block: Tag, data_url_prefix: str, preview_dir: Path) -> str:
    if block.name == "table":
        return _normalize_html_spacing(
            _render_html_node(block, data_url_prefix, preview_dir, prefer_inline_images=False)
        )
    return _paragraph_to_html(block, data_url_prefix, preview_dir)


def _paragraph_to_html(paragraph: Tag, data_url_prefix: str, preview_dir: Path) -> str:
    has_image = paragraph.find("img") is not None
    has_text = bool(paragraph.get_text(" ", strip=True))
    prefer_inline_images = has_image and has_text
    return _normalize_html_spacing(
        "".join(
            _render_html_node(
                child,
                data_url_prefix,
                preview_dir,
                prefer_inline_images=prefer_inline_images,
            )
            for child in paragraph.children
        )
    )


def _render_html_node(node, data_url_prefix: str, preview_dir: Path, *, prefer_inline_images: bool = False) -> str:
    if isinstance(node, NavigableString):
        text = str(node)
        return escape(text) if text else ""
    if not isinstance(node, Tag):
        return ""

    if node.name == "img":
        src = _normalize_image_src(node.get("src", ""), data_url_prefix, preview_dir)
        if not src:
            return ""
        alt = escape(node.get("alt", "公式图片"), quote=True)
        classes = " ".join(_image_classes(src, preview_dir, prefer_inline=prefer_inline_images))
        return f'<img class="{classes}" src="{escape(src, quote=True)}" alt="{alt}">'
    if node.name == "br":
        return "<br>"

    if node.name == "table":
        rows = "".join(
            _render_html_node(child, data_url_prefix, preview_dir, prefer_inline_images=False)
            for child in node.find_all("tr")
        )
        return f"<table>{rows}</table>" if rows else ""
    if node.name == "tr":
        cells = "".join(
            _render_html_node(child, data_url_prefix, preview_dir, prefer_inline_images=False)
            for child in node.find_all(["td", "th"], recursive=False)
        )
        return f"<tr>{cells}</tr>" if cells else ""
    if node.name in {"td", "th"}:
        children_html = "".join(
            _render_html_node(
                child,
                data_url_prefix,
                preview_dir,
                prefer_inline_images=True,
            )
            for child in node.children
        )
        tag_name = node.name
        return f"<{tag_name}>{children_html}</{tag_name}>"

    children_html = "".join(
        _render_html_node(
            child,
            data_url_prefix,
            preview_dir,
            prefer_inline_images=prefer_inline_images,
        )
        for child in node.children
    )
    if node.name in {"strong", "b", "em", "i", "sup", "sub"}:
        return f"<{node.name}>{children_html}</{node.name}>"
    if node.name == "span":
        classes = [item for item in (node.get("class") or []) if item]
        if classes:
            class_attr = escape(" ".join(classes), quote=True)
            return f'<span class="{class_attr}">{children_html}</span>'
        return f"<span>{children_html}</span>"
    return children_html


def _normalize_image_src(src: str, data_url_prefix: str, preview_dir: Path) -> str:
    src = (src or "").strip().replace("\\", "/")
    if not src:
        return ""
    if src.startswith(("http://", "https://", "/data/")):
        return src

    parsed = urlparse(src)
    if parsed.scheme == "file":
        src = unquote(parsed.path).replace("\\", "/")
    elif re.match(r"^[A-Za-z]:/", src):
        src = unquote(src)
    elif parsed.scheme and len(parsed.scheme) > 1:
        return src

    relative_src = _relative_preview_path(src, preview_dir)
    if relative_src.startswith("/"):
        return relative_src
    return data_url_prefix.rstrip("/") + "/" + relative_src.lstrip("./")


def _relative_preview_path(src: str, preview_dir: Path) -> str:
    normalized_src = unquote(src).replace("\\", "/")
    preview_abs = preview_dir.resolve()

    if re.match(r"^[A-Za-z]:/", normalized_src) or normalized_src.startswith("/"):
        try:
            return Path(normalized_src).resolve().relative_to(preview_abs).as_posix()
        except (OSError, ValueError):
            preview_text = preview_abs.as_posix().rstrip("/").lower()
            source_text = normalized_src.rstrip("/")
            if source_text.lower().startswith(preview_text + "/"):
                return source_text[len(preview_text) + 1 :]
            return normalized_src

    return normalized_src


def _normalize_html_spacing(fragment: str) -> str:
    fragment = re.sub(r"[ \t\r\n]+", " ", fragment)
    fragment = re.sub(r"\s+(</?br\s*/?>)", r"\1", fragment)
    return fragment.strip()


def _strip_question_number_from_html(fragment: str, number: str) -> str:
    pattern = (
        r"^\s*(?:[\(\uFF08\u3008\u3010]\s*[^\(\)\uFF08\uFF09\u3008\u3009\u3010\u30110-9]{1,6}\s*"
        r"[\)\uFF09\u3009\u3011]\s*)?"
        + re.escape(number)
        + r"\s*[\.\uFF0E\u3001]\s*"
    )
    stripped = re.sub(pattern, "", fragment, count=1).strip()
    if stripped != fragment.strip():
        return stripped
    return _strip_split_html_question_number(fragment, number)


def _strip_split_html_question_number(fragment: str, number: str) -> str:
    soup = BeautifulSoup(fragment, "html.parser")
    fullwidth = "".join(chr(ord(ch) + 0xFEE0) if "0" <= ch <= "9" else ch for ch in number)
    number_pattern = re.compile(rf"^\s*(?:{re.escape(number)}|{re.escape(fullwidth)})\s*$")
    punctuation_pattern = re.compile(r"^\s*[\.\uFF0E\u3001]\s*")
    saw_number = False

    for text_node in soup.find_all(string=True):
        value = str(text_node)
        if not saw_number:
            if not value.strip():
                continue
            if number_pattern.match(value):
                text_node.replace_with("")
                saw_number = True
                continue
            combined_pattern = re.compile(
                rf"^\s*(?:{re.escape(number)}|{re.escape(fullwidth)})\s*[\.\uFF0E\u3001]\s*"
            )
            updated = combined_pattern.sub("", value, count=1)
            if updated != value:
                text_node.replace_with(updated)
            break
        updated = punctuation_pattern.sub("", value, count=1)
        text_node.replace_with(updated)
        break

    return "".join(str(child) for child in soup.contents).strip()


def _strip_label_from_html(fragment: str, label: str) -> str:
    pattern = rf"^\s*[【\[]?\s*{label}\s*[】\]]?\s*[:：]?\s*"
    return re.sub(pattern, "", fragment, count=1).strip()


def _join_html_lines(lines: list[str]) -> str:
    return "".join(_wrap_html_line(line) for line in lines if line)


def _wrap_html_line(line: str) -> str:
    stripped = line.strip()
    if stripped.startswith("<table"):
        return stripped
    return f"<p>{line}</p>"


def _looks_like_option_group(line: str) -> bool:
    return len(re.findall(r"(?:^|\s)[A-HＡ-Ｈ]\s*[\.．、]", line)) >= 2


def _detect_section_type(line: str) -> str:
    compact = re.sub(r"\s+", "", line)
    header_like = bool(SECTION_PREFIX_RE.match(compact) or re.match(r"^.{0,12}题$", compact))
    if not header_like:
        return ""
    for keyword in SECTION_KEYWORDS:
        if keyword in compact:
            return keyword
    return ""


def _drop_preface_questions(questions: list[QuestionDraft]) -> list[QuestionDraft]:
    for index, question in enumerate(questions):
        if question.number != "1":
            continue
        suffix_numbers = [_safe_int(item.number) for item in questions[index:]]
        if suffix_numbers[:19] == list(range(1, 20)):
            return questions[index:]
    first_typed_index = next(
        (index for index, question in enumerate(questions) if question.question_type != "未分类"),
        None,
    )
    if first_typed_index and first_typed_index > 0:
        return questions[first_typed_index:]
    return questions


def _normalize_digits(text: str) -> str:
    normalized: list[str] = []
    for char in text:
        if char.isdigit():
            try:
                normalized.append(str(unicodedata.digit(char)))
                continue
            except (TypeError, ValueError):
                pass
        normalized.append(char)
    return "".join(normalized)


def _should_start_new_question(current_number: str | None, candidate_number: str) -> bool:
    if not candidate_number:
        return False
    if current_number is None:
        return True
    current_value = _safe_int(current_number)
    candidate_value = _safe_int(candidate_number)
    if current_value is None or candidate_value is None:
        return candidate_number != current_number
    return candidate_value > current_value


def _safe_int(value: str) -> int | None:
    if not value or not value.isdigit():
        return None
    return int(value)


def _image_classes(src: str, preview_dir: Path, *, prefer_inline: bool = False) -> list[str]:
    classes = ["inline-math"]
    if prefer_inline:
        classes.append("inline-formula")
        return classes
    local_path = _resolve_local_image_path(src, preview_dir)
    size = _read_local_image_size(local_path)
    if size and _looks_like_inline_formula_image(*size):
        classes.append("inline-formula")
    else:
        classes.append("block-figure")
    return classes


def _resolve_local_image_path(src: str, preview_dir: Path) -> Path | None:
    parsed = urlparse(src)
    relative = parsed.path if parsed.scheme or src.startswith("/") else src
    relative = relative.replace("\\", "/").lstrip("/")
    if relative.startswith("data/previews/"):
        preview_parts = preview_dir.parts
        if "data" not in preview_parts:
            remainder = Path(*Path(relative).parts[3:])
            return preview_dir / remainder
        data_root = Path(*preview_parts[: preview_parts.index("data") + 1])
        relative = relative.removeprefix("data/previews/")
        return data_root / "previews" / Path(relative)
    return preview_dir / Path(relative)


def _read_local_image_size(path: Path | None) -> tuple[int, int] | None:
    if path is None or not path.exists() or not path.is_file():
        return None
    try:
        with Image.open(path) as image:
            return image.size
    except Exception:
        return None


def _looks_like_inline_formula_image(width: int, height: int) -> bool:
    if width <= 0 or height <= 0:
        return False
    if height <= 120:
        return True
    ratio = width / height
    return height <= 220 and ratio >= 1.8


def _looks_like_solution(text: str) -> bool:
    if not text:
        return False
    text_lower = text.lower()
    for keyword in SOLUTION_KEYWORDS:
        if keyword in text_lower:
            return True
    return False


def _looks_like_sentence(text: str) -> bool:
    if not text:
        return False
    sentence_endings = ("。", "！", "？", ";", "；", "\n", "\r")
    for ending in sentence_endings:
        if ending in text:
            return True
    return len(text) > 15
